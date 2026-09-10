from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
import io, zipfile, re, base64, secrets, random, sqlite3, os
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
import qrcode
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A6
from reportlab.lib.utils import ImageReader

BASE = Path(__file__).resolve().parent
DB = BASE / "doorprize_pro.db"
app = Flask(__name__)
app.secret_key = os.environ.get("DOORPRIZE_SECRET_KEY", "doorprize-pro-local-secret-change-me")
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

def now():
    return datetime.now().isoformat(timespec="seconds")


def qr_png_bytes(token):
    qr = qrcode.QRCode(version=None, box_size=10, border=4)
    qr.add_data(token)
    qr.make(fit=True)
    img = qr.make_image()
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def qr_data_uri(token):
    return "data:image/png;base64," + base64.b64encode(qr_png_bytes(token).getvalue()).decode("ascii")


def ensure_schema(c):
    # Base tables first so migration also works with an old database.
    c.executescript("""
    CREATE TABLE IF NOT EXISTS settings (
      key TEXT PRIMARY KEY,
      value TEXT NOT NULL DEFAULT ''
    );
    CREATE TABLE IF NOT EXISTS participants (
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL,
      ticket TEXT DEFAULT '',
      department TEXT DEFAULT '',
      phone TEXT DEFAULT '',
      is_winner INTEGER NOT NULL DEFAULT 0,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS prizes (
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL,
      quantity INTEGER NOT NULL DEFAULT 1,
      remaining INTEGER NOT NULL DEFAULT 1,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS draws (
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      participant_id INTEGER NOT NULL,
      prize_id INTEGER,
      drawn_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS attendance_logs (
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      participant_id INTEGER NOT NULL,
      scanned_at TEXT NOT NULL
    );
    """)
    cols = {r[1] for r in c.execute("PRAGMA table_info(participants)").fetchall()}
    if "qr_token" not in cols:
        c.execute("ALTER TABLE participants ADD COLUMN qr_token TEXT")
    if "attendance_status" not in cols:
        c.execute("ALTER TABLE participants ADD COLUMN attendance_status TEXT NOT NULL DEFAULT 'BELUM_HADIR'")
    if "checked_in_at" not in cols:
        c.execute("ALTER TABLE participants ADD COLUMN checked_in_at TEXT")

    defaults = {
        "event_name": "GRAND DOORPRIZE",
        "event_subtitle": "Konfirmasi Kehadiran & Undian",
        "logo_data": "",
        "theme": "midnight",
    }
    for k, v in defaults.items():
        c.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))
    ensure_admin_account(c)

    # Repair old records which were imported before QR attendance was added.
    missing = c.execute("SELECT id FROM participants WHERE qr_token IS NULL OR TRIM(qr_token)='' ").fetchall()
    for row in missing:
        c.execute("UPDATE participants SET qr_token=? WHERE id=?", (secrets.token_urlsafe(24), row[0]))
    # Normalize null/empty attendance values.
    c.execute("UPDATE participants SET attendance_status='BELUM_HADIR' WHERE attendance_status IS NULL OR TRIM(attendance_status)='' ")
    c.commit()


def get_db():
    c = sqlite3.connect(DB, timeout=10)
    c.row_factory = sqlite3.Row
    ensure_schema(c)
    return c


def settings_dict():
    c = get_db()
    rows = c.execute("SELECT key,value FROM settings").fetchall()
    c.close()
    return {r["key"]: r["value"] for r in rows}


def is_admin():
    return bool(session.get("admin_logged_in"))

def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if is_admin():
            return view(*args, **kwargs)
        if request.path.startswith("/api/"):
            return jsonify(success=False, error="Akses admin diperlukan. Silakan login."), 401
        return redirect(url_for("admin_login", next=request.path))
    return wrapped

def ensure_admin_account(c):
    # Default local admin account. The password is stored as a secure hash in SQLite.
    row=c.execute("SELECT value FROM settings WHERE key='admin_password_hash'").fetchone()
    if not row or not row[0]:
        c.execute("INSERT INTO settings(key,value) VALUES('admin_username','admin') ON CONFLICT(key) DO NOTHING")
        c.execute("INSERT INTO settings(key,value) VALUES('admin_password_hash',?) ON CONFLICT(key) DO NOTHING", (generate_password_hash("admin123"),))
        c.commit()

@app.errorhandler(413)
def too_large(e):
    if request.path.startswith("/api/"):
        return jsonify(error="File terlalu besar. Maksimal 10 MB."), 413
    return "File terlalu besar. Maksimal 10 MB.", 413


@app.errorhandler(Exception)
def handle_exception(e):
    app.logger.exception("Unhandled error")
    if request.path.startswith("/api/"):
        return jsonify(error=f"Server error: {e}"), 500
    raise e

@app.post("/api/settings")
@admin_required
def save_settings():
    data=request.json or {}
    c=get_db()
    for k in ("event_name","event_subtitle","theme"):
        if k in data:
            c.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(k,str(data[k])))
    c.commit(); c.close()
    return jsonify(settings_dict())

@app.post("/api/logo")
@admin_required
def upload_logo():
    f=request.files.get("logo")
    if not f or not f.filename:
        return jsonify({"error":"File logo tidak ditemukan"}),400
    if f.mimetype not in ("image/png","image/jpeg","image/webp"):
        return jsonify({"error":"Logo harus PNG, JPG, atau WEBP"}),400
    raw=f.read()
    data="data:"+f.mimetype+";base64,"+base64.b64encode(raw).decode()
    c=get_db(); c.execute("INSERT INTO settings(key,value) VALUES('logo_data',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(data,))
    c.commit(); c.close()
    return jsonify({"ok":True,"logo_data":data})


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if is_admin():
        return redirect(request.args.get("next") or url_for("register_and_setting_page"))
    error = ""
    next_url = request.args.get("next") or request.form.get("next") or ""
    if request.method == "POST":
        username = str(request.form.get("username") or "").strip()
        password = str(request.form.get("password") or "")
        c = get_db()
        user = c.execute("SELECT value FROM settings WHERE key='admin_username'").fetchone()
        ph = c.execute("SELECT value FROM settings WHERE key='admin_password_hash'").fetchone()
        c.close()
        if user and ph and username == user[0] and check_password_hash(ph[0], password):
            session.clear()
            session["admin_logged_in"] = True
            session["admin_username"] = username
            return redirect(next_url if next_url.startswith("/") else url_for("register_and_setting_page"))
        error = "Username atau password admin salah."
    return render_template("admin_login.html", error=error, next=next_url)

@app.get("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin_login"))

@app.post("/admin/change-password")
@admin_required
def admin_change_password():
    current = str(request.form.get("current_password") or "")
    new_password = str(request.form.get("new_password") or "")
    confirm = str(request.form.get("confirm_password") or "")

    if len(new_password) < 8:
        return redirect(url_for("register_and_setting_page", password_error="Password baru minimal 8 karakter."))
    if new_password != confirm:
        return redirect(url_for("register_and_setting_page", password_error="Konfirmasi password tidak sama."))

    c = get_db()
    row = c.execute("SELECT value FROM settings WHERE key='admin_password_hash'").fetchone()
    if not row or not check_password_hash(row[0], current):
        c.close()
        return redirect(url_for("register_and_setting_page", password_error="Password lama salah."))

    c.execute("INSERT INTO settings(key,value) VALUES('admin_password_hash',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
              (generate_password_hash(new_password),))
    c.commit()
    c.close()
    return redirect(url_for("register_and_setting_page", password_changed="1"))


@app.route("/")
def index():
    return render_template("index.html", settings=settings_dict())

@app.route("/register_and_setting")
@admin_required
def register_and_setting_page():
    return render_template("register_and_setting.html", settings=settings_dict())

@app.route("/attendance")
@app.route("/konfirmasi-kehadiran")
def attendance_page():
    return render_template("attendance.html", settings=settings_dict())

@app.route("/state")
def state_page():
    return render_template("state.html", settings=settings_dict())

@app.route("/draw")
def draw_page():
    return render_template("draw.html", settings=settings_dict())


@app.route("/qr-cards")
def qr_cards():
    c = get_db()
    rows = c.execute("SELECT id,name,ticket,department,phone,qr_token,attendance_status,checked_in_at FROM participants ORDER BY name COLLATE NOCASE").fetchall()
    participants = []
    for r in rows:
        item = dict(r)
        if not item.get("qr_token"):
            item["qr_token"] = secrets.token_urlsafe(24)
            c.execute("UPDATE participants SET qr_token=? WHERE id=?", (item["qr_token"], item["id"]))
        item["qr_image"] = qr_data_uri(item["qr_token"])
        participants.append(item)
    c.commit(); c.close()
    return render_template("qr_cards.html", participants=participants, settings=settings_dict())


@app.get("/api/state")
def state():
    c = get_db()
    participants = [dict(r) for r in c.execute("SELECT * FROM participants ORDER BY id DESC")]
    prizes = [dict(r) for r in c.execute("SELECT * FROM prizes ORDER BY id DESC")]
    winners = [dict(r) for r in c.execute("""
      SELECT d.id,d.drawn_at,p.name,p.ticket,p.department,pr.name AS prize
      FROM draws d JOIN participants p ON p.id=d.participant_id
      LEFT JOIN prizes pr ON pr.id=d.prize_id ORDER BY d.id DESC
    """)]
    stats = {
        "participants": len(participants),
        "eligible": c.execute("SELECT COUNT(*) FROM participants WHERE is_winner=0 AND attendance_status='HADIR'").fetchone()[0],
        "winners": c.execute("SELECT COUNT(*) FROM participants WHERE is_winner=1").fetchone()[0],
        "hadir": c.execute("SELECT COUNT(*) FROM participants WHERE attendance_status='HADIR'").fetchone()[0],
        "belum_hadir": c.execute("SELECT COUNT(*) FROM participants WHERE attendance_status!='HADIR'").fetchone()[0],
        "prizes": sum(p["quantity"] for p in prizes),
        "remaining_prizes": sum(p["remaining"] for p in prizes),
        "draws": len(winners),
    }
    c.close()
    return jsonify(settings=settings_dict(), participants=participants, prizes=prizes, winners=winners, stats=stats)


@app.get("/api/attendance")
def attendance():
    c = get_db()
    rows = c.execute("""
      SELECT id,name,ticket,department,phone,qr_token,attendance_status,checked_in_at,is_winner
      FROM participants ORDER BY name COLLATE NOCASE
    """).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])


def find_participant(c, token="", ticket=""):
    token = token.strip(); ticket = ticket.strip()
    if token:
        row = c.execute("SELECT * FROM participants WHERE qr_token=?", (token,)).fetchone()
        if row:
            return row
    if ticket:
        # Case-insensitive ticket lookup, useful when Excel stores text/number inconsistently.
        row = c.execute("SELECT * FROM participants WHERE TRIM(ticket)=TRIM(?) COLLATE NOCASE", (ticket,)).fetchone()
        if row:
            return row
    return None


@app.post("/api/attendance/checkin")
def attendance_checkin():
    data = request.get_json(silent=True) or {}
    token = str(data.get("token") or "").strip()
    ticket = str(data.get("ticket") or "").strip()
    if not token and not ticket:
        return jsonify(success=False, error="QR Code atau nomor tiket kosong."), 400

    c = get_db()
    row = find_participant(c, token, ticket)
    if not row:
        c.close()
        return jsonify(success=False, error="QR/nomor tiket tidak ditemukan."), 404

    if str(row["attendance_status"] or "").upper() == "HADIR":
        participant = dict(row)
        c.close()
        return jsonify(success=True, already=True, participant=participant,
                       message="Peserta sudah terkonfirmasi hadir.")

    checked = now()
    c.execute("UPDATE participants SET attendance_status='HADIR', checked_in_at=? WHERE id=?", (checked, row["id"]))
    c.execute("INSERT INTO attendance_logs(participant_id,scanned_at) VALUES(?,?)", (row["id"], checked))
    c.commit()
    updated = c.execute("SELECT * FROM participants WHERE id=?", (row["id"],)).fetchone()
    c.close()
    return jsonify(success=True, already=False, participant=dict(updated),
                   message="Kehadiran berhasil dikonfirmasi.")


@app.post("/api/attendance/reset")
def attendance_reset():
    c = get_db()
    c.execute("UPDATE participants SET attendance_status='BELUM_HADIR', checked_in_at=NULL")
    c.execute("DELETE FROM attendance_logs")
    c.commit(); c.close()
    return jsonify(success=True, message="Semua status kehadiran dikembalikan menjadi BELUM HADIR.")


@app.post("/api/participants")
@admin_required
def add_participant():
    d = request.get_json(silent=True) or {}
    name = str(d.get("name") or "").strip()
    ticket = str(d.get("ticket") or "").strip()
    if not ticket:
        return jsonify(error="Nomor Tiket wajib diisi."), 400
    c = get_db()
    if c.execute("SELECT 1 FROM participants WHERE TRIM(ticket)=TRIM(?) COLLATE NOCASE", (ticket,)).fetchone():
        c.close(); return jsonify(error="Nomor Tiket sudah terdaftar."), 400
    cur = c.execute("""INSERT INTO participants
      (name,ticket,department,phone,is_winner,created_at,qr_token,attendance_status)
      VALUES(?,?,?,?,?,?,?,?)""", (name, ticket, str(d.get("department") or ""),
      str(d.get("phone") or ""), 0, now(), secrets.token_urlsafe(24), "BELUM_HADIR"))
    c.commit(); row = c.execute("SELECT * FROM participants WHERE id=?", (cur.lastrowid,)).fetchone(); c.close()
    return jsonify(dict(row))

@app.delete("/api/participants/<int:pid>")
@admin_required
def del_participant(pid):
    c=get_db(); c.execute("DELETE FROM participants WHERE id=?",(pid,)); c.commit(); c.close()
    return jsonify({"ok":True})

@app.post("/api/import-excel")
@admin_required
def import_excel():
    f = request.files.get("file")
    filename = (f.filename or "").strip().lower() if f else ""
    if not f or not filename.endswith((".xlsx", ".xlsm")):
        return jsonify(error="Unggah file Excel .xlsx atau .xlsm"), 400

    wb = None
    c = None
    try:
        raw = f.read()
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify(error="Sheet Excel kosong"), 400

        # Nomor Tiket = WAJIB. Nama Peserta = OPSIONAL.
        aliases = {
            "no": "_no", "nomor": "_no", "no.": "_no",
            "nama": "name", "name": "name", "peserta": "name", "nama peserta": "name",
            "ticket": "ticket", "tiket": "ticket", "nomor tiket": "ticket", "no tiket": "ticket", "no. tiket": "ticket",
            "departemen": "department", "department": "department", "bagian": "department", "divisi": "department", "unit kerja": "department",
            "phone": "phone", "telepon": "phone", "no hp": "phone", "no. hp": "phone", "nomor hp": "phone",
            "nomor hp/wa": "phone", "no hp/wa": "phone", "whatsapp": "phone", "wa": "phone"
        }
        headers = [str(x or "").strip().lower() for x in rows[0]]
        mapped = [aliases.get(h, h) for h in headers]

        if "ticket" not in mapped:
            return jsonify(error="Kolom Nomor Tiket wajib ada. Gunakan header: Nomor Tiket."), 400

        c = get_db()
        count = 0
        skipped = 0
        skipped_details = []
        seen_tickets = set()

        def clean_ticket(value):
            if value is None:
                return ""
            # Excel dapat membaca nomor tiket numerik sebagai int/float.
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            return str(value).strip()

        for row_no, vals in enumerate(rows[1:], start=2):
            item = {mapped[i]: (vals[i] if i < len(vals) else "") for i in range(len(mapped))}
            ticket = clean_ticket(item.get("ticket"))
            name = str(item.get("name") or "").strip()
            department = str(item.get("department") or "").strip()
            phone = str(item.get("phone") or "").strip()

            if not ticket:
                skipped += 1
                skipped_details.append(f"Baris {row_no}: Nomor Tiket kosong")
                continue

            ticket_key = ticket.casefold()
            if ticket_key in seen_tickets:
                skipped += 1
                skipped_details.append(f"Baris {row_no}: Nomor Tiket duplikat di file ({ticket})")
                continue

            if c.execute("SELECT 1 FROM participants WHERE TRIM(ticket)=TRIM(?) COLLATE NOCASE", (ticket,)).fetchone():
                skipped += 1
                skipped_details.append(f"Baris {row_no}: Nomor Tiket sudah terdaftar ({ticket})")
                continue

            seen_tickets.add(ticket_key)
            c.execute("""INSERT INTO participants
              (name,ticket,department,phone,is_winner,created_at,qr_token,attendance_status)
              VALUES(?,?,?,?,?,?,?,?)""", (
                name, ticket, department, phone, 0, now(),
                secrets.token_urlsafe(24), "BELUM_HADIR"
            ))
            count += 1

        c.commit()
        return jsonify(
            inserted=count,
            skipped=skipped,
            skipped_details=skipped_details[:100],
            sheet=ws.title,
            message=f"{count} peserta berhasil diimpor. {skipped} baris dilewati karena Nomor Tiket kosong/duplikat."
        )
    except Exception as e:
        if c:
            try:
                c.rollback()
            except Exception:
                pass
        return jsonify(error=f"Gagal membaca Excel: {e}"), 400
    finally:
        if c:
            c.close()
        if wb:
            wb.close()


@app.post("/api/import-prizes")
@admin_required
def import_prizes():
    f = request.files.get("file")
    filename = (f.filename or "").strip().lower() if f else ""
    if not f or not filename.endswith((".xlsx", ".xlsm")):
        return jsonify(error="Unggah file Excel hadiah .xlsx atau .xlsm"), 400
    try:
        raw = f.read(); wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True); ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows: wb.close(); return jsonify(error="Sheet Excel hadiah kosong"), 400
        aliases = {"nama":"name","name":"name","hadiah":"name","prize":"name","nama hadiah":"name",
                   "jumlah":"quantity","qty":"quantity","quantity":"quantity","stok":"quantity"}
        headers = [str(x or "").strip().lower() for x in rows[0]]; mapped = [aliases.get(h,h) for h in headers]
        if "name" not in mapped:
            wb.close(); return jsonify(error="Kolom nama hadiah tidak ditemukan."), 400
        c = get_db(); count = units = 0
        for vals in rows[1:]:
            item = {mapped[i]: (vals[i] if i < len(vals) else "") for i in range(len(mapped))}
            name = str(item.get("name") or "").strip()
            if not name: continue
            try: qty = max(1, int(float(item.get("quantity", 1) or 1)))
            except: qty = 1
            c.execute("INSERT INTO prizes(name,quantity,remaining,created_at) VALUES(?,?,?,?)", (name,qty,qty,now()))
            count += 1; units += qty
        c.commit(); c.close(); wb.close()
        return jsonify(inserted=count, units=units, sheet=ws.title, message=f"{count} hadiah berhasil diimpor ({units} stok)")
    except Exception as e:
        return jsonify(error=f"Gagal membaca Excel hadiah: {e}"), 400


@app.get("/api/qr/<int:participant_id>.png")
def participant_qr(participant_id):
    c = get_db(); row = c.execute("SELECT * FROM participants WHERE id=?", (participant_id,)).fetchone()
    if not row:
        c.close(); return jsonify(error="Peserta tidak ditemukan."), 404
    token = row["qr_token"] or secrets.token_urlsafe(24)
    if not row["qr_token"]:
        c.execute("UPDATE participants SET qr_token=? WHERE id=?", (token, participant_id)); c.commit()
    c.close()
    return send_file(qr_png_bytes(token), mimetype="image/png", as_attachment=False)


def _safe_filename(value, fallback="kartu"):
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "").strip())
    return value.strip("._") or fallback


def participant_card_pdf_bytes(row, settings):
    """Generate one printable participant card PDF, matching the QR card content."""
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=A6)
    width, height = A6

    # Premium card frame/background
    c.setFillColorRGB(0.055, 0.106, 0.208)
    c.roundRect(10, 10, width - 20, height - 20, 14, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.roundRect(18, 18, width - 36, height - 36, 10, fill=1, stroke=0)

    # Event logo (optional, stored as data URI in settings)
    logo = settings.get("logo_data", "") if settings else ""
    if logo.startswith("data:image") and "," in logo:
        try:
            raw = base64.b64decode(logo.split(",", 1)[1])
            c.drawImage(ImageReader(io.BytesIO(raw)), 22, height - 62, width=44, height=34,
                        preserveAspectRatio=True, anchor='c', mask='auto')
        except Exception:
            pass

    event_name = settings.get("event_name") or settings.get("title") or "DOORPRIZE PRO"
    c.setFillColorRGB(0.055, 0.106, 0.208)
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString(width / 2, height - 45, str(event_name)[:34])

    name = str(row["name"] or "Peserta")
    ticket = str(row["ticket"] or "-")
    department = str(row["department"] or "-")

    c.setFont("Helvetica-Bold", 17)
    # Fit long names into card width
    font_size = 17
    while font_size > 10 and c.stringWidth(name, "Helvetica-Bold", font_size) > width - 50:
        font_size -= 1
    c.setFont("Helvetica-Bold", font_size)
    c.drawCentredString(width / 2, height - 88, name[:45])

    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(width / 2, height - 108, f"No. Tiket: {ticket}")
    c.setFont("Helvetica", 10)
    c.drawCentredString(width / 2, height - 125, department[:34])

    token = row["qr_token"] or secrets.token_urlsafe(24)
    qr = qrcode.QRCode(version=None, box_size=8, border=3)
    qr.add_data(token); qr.make(fit=True)
    img = qr.make_image()
    qrbuf = io.BytesIO(); img.save(qrbuf, format="PNG"); qrbuf.seek(0)
    qr_size = 135
    c.drawImage(ImageReader(qrbuf), (width - qr_size) / 2, 62, width=qr_size, height=qr_size,
                preserveAspectRatio=True, mask='auto')

    c.setFillColorRGB(0.35, 0.40, 0.50)
    c.setFont("Helvetica", 8)
    c.drawCentredString(width / 2, 45, "Scan QR untuk konfirmasi kehadiran")
    c.showPage(); c.save(); out.seek(0)
    return out


@app.get("/api/qr/<int:participant_id>.pdf")
def download_participant_qr(participant_id):
    db = get_db()
    row = db.execute("SELECT id,name,ticket,department,qr_token FROM participants WHERE id=?", (participant_id,)).fetchone()
    if not row:
        db.close(); return jsonify({"error": "Peserta tidak ditemukan."}), 404
    token = row["qr_token"]
    if not token:
        token = secrets.token_urlsafe(24)
        db.execute("UPDATE participants SET qr_token=? WHERE id=?", (token, participant_id)); db.commit()
        row = db.execute("SELECT id,name,ticket,department,qr_token FROM participants WHERE id=?", (participant_id,)).fetchone()
    settings = settings_dict()
    pdf = participant_card_pdf_bytes(row, settings)
    db.close()
    filename = f"{_safe_filename(row['ticket'], str(row['id']))}.pdf"
    return send_file(pdf, mimetype="application/pdf", as_attachment=True, download_name=filename)


@app.get("/api/qr/all.zip")
def all_qr():
    c = get_db(); rows = c.execute("SELECT id,name,ticket,department,qr_token FROM participants ORDER BY name COLLATE NOCASE").fetchall()
    if not rows:
        c.close(); return jsonify(error="Belum ada peserta."), 404
    settings = settings_dict()
    bundle = io.BytesIO()
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        for row in rows:
            token = row["qr_token"] or secrets.token_urlsafe(24)
            if not row["qr_token"]:
                c.execute("UPDATE participants SET qr_token=? WHERE id=?", (token, row["id"]))
                row = dict(row); row["qr_token"] = token
            pdf = participant_card_pdf_bytes(row, settings).getvalue()
            ticket = _safe_filename(row["ticket"], str(row["id"]))
            z.writestr(f"{ticket}.pdf", pdf)
    c.commit(); c.close(); bundle.seek(0)
    return send_file(bundle, mimetype="application/zip", as_attachment=True, download_name="Kartu_Peserta_Doorprize.zip")

@app.post("/api/prizes")
@admin_required
def add_prize():
    d=request.json or {}; name=str(d.get("name","")).strip()
    try:q=max(1,int(d.get("quantity",1)))
    except:q=1
    if not name:return jsonify({"error":"Nama hadiah wajib diisi"}),400
    c=get_db(); cur=c.execute("INSERT INTO prizes(name,quantity,remaining,created_at) VALUES(?,?,?,?)",(name,q,q,datetime.now().isoformat(timespec="seconds")))
    c.commit(); row=c.execute("SELECT * FROM prizes WHERE id=?",(cur.lastrowid,)).fetchone(); c.close()
    return jsonify(dict(row))

@app.delete("/api/prizes/<int:pid>")
@admin_required
def del_prize(pid):
    c=get_db(); c.execute("DELETE FROM prizes WHERE id=?",(pid,)); c.commit(); c.close(); return jsonify({"ok":True})


@app.post("/api/draw")
def draw():
    d = request.get_json(silent=True) or {}; prize_id = d.get("prize_id")
    c = get_db(); prize = None
    if prize_id:
        prize = c.execute("SELECT * FROM prizes WHERE id=?", (prize_id,)).fetchone()
        if not prize or prize["remaining"] <= 0:
            c.close(); return jsonify(error="Hadiah habis"), 400
    candidates = c.execute("SELECT * FROM participants WHERE is_winner=0 AND attendance_status='HADIR'").fetchall()
    if not candidates:
        c.close(); return jsonify(error="Tidak ada peserta HADIR yang eligible."), 400
    winner = random.choice(candidates); ts = now()
    c.execute("UPDATE participants SET is_winner=1 WHERE id=?", (winner["id"],))
    if prize: c.execute("UPDATE prizes SET remaining=remaining-1 WHERE id=?", (prize["id"],))
    cur = c.execute("INSERT INTO draws(participant_id,prize_id,drawn_at) VALUES(?,?,?)", (winner["id"], prize["id"] if prize else None, ts))
    c.commit(); result = {"draw_id":cur.lastrowid,"participant":dict(winner),"prize":dict(prize) if prize else None,"drawn_at":ts}; c.close()
    return jsonify(result)


@app.post("/api/reset")
def reset():
    c=get_db(); c.execute("DELETE FROM draws"); c.execute("UPDATE participants SET is_winner=0"); c.execute("UPDATE prizes SET remaining=quantity"); c.commit(); c.close(); return jsonify(ok=True)

@app.post("/api/reset-all")
@admin_required
def reset_all():
    c=get_db(); c.execute("DELETE FROM draws"); c.execute("DELETE FROM participants")
    c.execute("DELETE FROM prizes"); c.commit(); c.close(); return jsonify({"ok":True})

if __name__ == "__main__":
    # Debug=False prevents browser-facing tracebacks during the event.
    app.run(host="0.0.0.0", port=5000, debug=False)